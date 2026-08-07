#!/usr/bin/env python3
"""Build docs/plan-sponsor-pitch.pptx from the plan sponsor pitch content."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "plan-sponsor-pitch.pptx"

INK = RGBColor(0x12, 0x16, 0x1F)
INK_700 = RGBColor(0x40, 0x4B, 0x60)
INK_600 = RGBColor(0x4E, 0x5D, 0x76)
INK_200 = RGBColor(0xD5, 0xDA, 0xE3)
INK_50 = RGBColor(0xF6, 0xF7, 0xF9)
GLASS = RGBColor(0x1C, 0xA2, 0xA7)
GLASS_DARK = RGBColor(0x14, 0x81, 0x87)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF4, 0xF6, 0xF8)
AMBER = RGBColor(0xB4, 0x53, 0x09)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    return shape


def add_round_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    # Less pill-like corners
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def para(tf, text, size=14, bold=False, color=INK, space_after=6, font="Calibri", align=None):
    p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
    if tf.text and p == tf.paragraphs[0] and tf.paragraphs[0].text == "":
        # first paragraph already empty
        pass
    elif tf.text:
        p = tf.add_paragraph()
    p.clear() if hasattr(p, "clear") else None
    if not tf.text and tf.paragraphs[0].runs:
        # replace default
        run = tf.paragraphs[0].runs[0]
        run.text = text
        set_run(run, size, bold, color, font)
        p = tf.paragraphs[0]
    else:
        if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1 and not any(r.text for r in tf.paragraphs[0].runs):
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = text
            set_run(run, size, bold, color, font)
        else:
            p = tf.add_paragraph()
            run = p.add_run()
            run.text = text
            set_run(run, size, bold, color, font)
    p.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def write_block(shape, lines):
    """lines: list of (text, size, bold, color, space_after)."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, (text, size, bold, color, space_after) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0:
            p.clear()
            # after clear, ensure a run
        run = p.add_run()
        run.text = text
        set_run(run, size, bold, color)
        p.space_after = Pt(space_after)


def eyebrow(slide, text, top=Inches(0.45), dark=False):
    box = add_textbox(slide, Inches(0.7), top, Inches(12), Inches(0.35))
    write_block(
        box,
        [(text.upper(), 11, True, GLASS if not dark else RGBColor(0x74, 0xD8, 0xD8), 0)],
    )


def title(slide, text, top=Inches(0.75), dark=False, width=Inches(11)):
    box = add_textbox(slide, Inches(0.7), top, width, Inches(1.3))
    write_block(box, [(text, 36, False, WHITE if dark else INK, 0,)])
    # Georgia for display feel
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = "Georgia"


def lede(slide, text, top=Inches(2.0), dark=False):
    box = add_textbox(slide, Inches(0.7), top, Inches(11.5), Inches(0.9))
    write_block(
        box,
        [(text, 15, False, RGBColor(0xA8, 0xB0, 0xBE) if dark else INK_600, 0)],
    )


def card(slide, left, top, width, height, heading, bullets, dark=False):
    bg = RGBColor(0x1D, 0x24, 0x32) if dark else WHITE
    if not dark:
        add_round_rect(slide, left, top, width, height, WHITE)
        border = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        border.fill.background()
        border.line.color.rgb = INK_200
        try:
            border.adjustments[0] = 0.08
        except Exception:
            pass
    else:
        add_round_rect(slide, left, top, width, height, bg)

    box = add_textbox(slide, left + Inches(0.2), top + Inches(0.15), width - Inches(0.35), height - Inches(0.25))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = heading
    set_run(r, 14, True, WHITE if dark else INK)
    p.space_after = Pt(8)
    for b in bullets:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + b
        set_run(r, 12, False, RGBColor(0xC5, 0xCB, 0xD4) if dark else INK_700)
        p.space_after = Pt(4)


def stat_card(slide, left, top, width, height, value, label, dark=False):
    if dark:
        add_round_rect(slide, left, top, width, height, RGBColor(0x1D, 0x24, 0x32))
    else:
        add_round_rect(slide, left, top, width, height, WHITE)
        border = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        border.fill.background()
        border.line.color.rgb = INK_200
        try:
            border.adjustments[0] = 0.08
        except Exception:
            pass
    box = add_textbox(slide, left + Inches(0.2), top + Inches(0.2), width - Inches(0.35), height - Inches(0.3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = value
    set_run(r, 24, False, WHITE if dark else INK)
    r.font.name = "Georgia"
    p.space_after = Pt(6)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = label
    set_run(r, 12, False, RGBColor(0xA8, 0xB0, 0xBE) if dark else INK_600)


def callout(slide, text, title="Note", top=Inches(6.35), warn=False):
    bar = add_rect(slide, Inches(0.7), top, Inches(0.08), Inches(0.7), AMBER if warn else GLASS)
    bg = add_rect(
        slide,
        Inches(0.78),
        top,
        Inches(11.8),
        Inches(0.7),
        RGBColor(0xFF, 0xF7, 0xED) if warn else RGBColor(0xEE, 0xFB, 0xFA),
    )
    box = add_textbox(slide, Inches(1.0), top + Inches(0.08), Inches(11.3), Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title + "  "
    set_run(r, 11, True, INK_700)
    r = p.add_run()
    r.text = text
    set_run(r, 12, False, INK_700)


def table_slide(slide, headers, rows, top=Inches(2.85), col_widths=None):
    cols = len(headers)
    table_shape = slide.shapes.add_table(
        len(rows) + 1, cols, Inches(0.7), top, Inches(11.9), Inches(0.38 * (len(rows) + 1))
    )
    table = table_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_run(r, 10, True, INK_600)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK_50
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    set_run(r, 11, ci == 0, INK if ci == 0 else INK_700)


def blank_slide(prs, dark=False):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, INK if dark else PAPER)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 Title
    s = blank_slide(prs, dark=True)
    mark = add_round_rect(s, Inches(0.7), Inches(0.55), Inches(0.38), Inches(0.38), GLASS)
    name = add_textbox(s, Inches(1.2), Inches(0.55), Inches(3), Inches(0.4))
    write_block(name, [("Glass", 16, True, WHITE, 0)])
    eyebrow(s, "Transparent pharmacy benefit management", top=Inches(1.2), dark=True)
    title(s, "Glass for plan sponsors", top=Inches(1.55), dark=True)
    lede(
        s,
        "Best-in-class reporting, radical transparency without locking you into one fee model, and member service that answers in chat and text — from the same rules that adjudicate the claim.",
        top=Inches(3.0),
        dark=True,
    )
    w = Inches(3.85)
    stat_card(s, Inches(0.7), Inches(4.2), w, Inches(1.5), "Reporting", "Contract questions answered from the claim ledger", dark=True)
    stat_card(s, Inches(4.75), Inches(4.2), w, Inches(1.5), "Transparency", "Every dollar named — any commercial model", dark=True)
    stat_card(s, Inches(8.8), Inches(4.2), w, Inches(1.5), "Service", "Chat and text, grounded in live plan data", dark=True)
    box = add_textbox(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.9))
    write_block(
        box,
        [
            (
                "How we lead — Reports finance trusts, economics you can audit, and service that does not make people wait on hold. Flexible pricing is available — opacity is not.",
                13,
                False,
                RGBColor(0xC5, 0xCB, 0xD4),
                0,
            )
        ],
    )

    # 2 Problem
    s = blank_slide(prs)
    eyebrow(s, "The status quo")
    title(s, "What sponsors are tired of")
    lede(
        s,
        "The pharmacy benefit often fails in three places that matter more than the brochure discount: reports that cannot be checked, fees that hide inside drug cost, and member service that cannot explain a reject.",
        top=Inches(2.05),
    )
    card(
        s,
        Inches(0.7),
        Inches(3.1),
        Inches(3.85),
        Inches(2.9),
        "Reporting today",
        [
            "Quarterly PDFs with stored totals",
            "Guarantee credits that never quite appear",
            "Trend stories without claim-level proof",
            "Finance and benefits see different numbers",
        ],
    )
    card(
        s,
        Inches(4.75),
        Inches(3.1),
        Inches(3.85),
        Inches(2.9),
        "Transparency today",
        [
            "“Transparency” as a slide, not a ledger",
            "Pass-through sold, then exceptions",
            "Or spread accepted because RFPs demand PEPM",
            "No way to name every dollar either way",
        ],
    )
    card(
        s,
        Inches(8.8),
        Inches(3.1),
        Inches(3.85),
        Inches(2.9),
        "Service today",
        [
            "Phone trees and callback queues",
            "Chatbots that cannot read the claim",
            "Members escalate what should be one text",
            "HR becomes the unofficial help desk",
        ],
    )

    # 3 Pillars
    s = blank_slide(prs)
    eyebrow(s, "The offer")
    title(s, "Three pillars")
    lede(
        s,
        "Everything else — pass-through vs PEPM, shared savings, clinical modules — sits under these three.",
        top=Inches(2.05),
    )
    pillars = [
        (
            "1",
            "Best-in-class reporting",
            "Sponsor dashboards, guarantee reconciliation, rebate waterfalls, trend drivers, design scenarios, and a data agent for ad-hoc questions — all computed from claims. Finance can drill to the dollar.",
        ),
        (
            "2",
            "Transparency with flexible pricing",
            "Claim-level derivation stays mandatory. How Glass gets paid is a menu: pass-through, PEPM, shared savings, outcomes at risk, or disclosed margin — always as named lines.",
        ),
        (
            "3",
            "Excellent service by chat and text",
            "Members get answers where they already are — chat and SMS — from an agent grounded in live eligibility, formulary, PA status, and claim history. Every figure comes from the book.",
        ),
    ]
    y = Inches(3.0)
    for n, h, body in pillars:
        add_round_rect(s, Inches(0.7), y, Inches(0.45), Inches(0.45), INK)
        num = add_textbox(s, Inches(0.7), y + Inches(0.05), Inches(0.45), Inches(0.4))
        write_block(num, [(n, 16, True, WHITE, 0)])
        for p in num.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
        box = add_textbox(s, Inches(1.35), y - Inches(0.05), Inches(11.2), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = h
        set_run(r, 16, True, INK)
        p.space_after = Pt(4)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = body
        set_run(r, 13, False, INK_600)
        y += Inches(1.25)

    # 4 Reporting
    s = blank_slide(prs)
    eyebrow(s, "Pillar one")
    title(s, "Best-in-class reporting")
    lede(
        s,
        "The questions a plan sponsor should answer at year-end — computed from the claim ledger, not supplied by the PBM as a finished answer.",
        top=Inches(2.05),
    )
    table_slide(
        s,
        ["Report", "What it answers", "Why it is best-in-class"],
        [
            ["Guarantee reconciliation", "Did pricing and ops promises pay?", "Each measure settles alone — no cross-category netting"],
            ["Rebate waterfall", "What was earned vs retained vs passed?", "Floors, admin, and surplus as separate lines"],
            ["Spread / fee comparison", "What would the same claims cost elsewhere?", "Identical book, alternate economics — a calculation"],
            ["Trend drivers", "Why did PMPM move?", "Price, util, mix, starts, leavers that sum to the whole"],
            ["Change scenarios", "What if we change the design?", "Whole-book re-adjudication + named member disruption"],
            ["Data agent", "Ad-hoc questions and report briefings", "Ask in plain language; every figure is tool-backed"],
        ],
        top=Inches(2.95),
        col_widths=[Inches(3.0), Inches(4.2), Inches(4.7)],
    )
    callout(
        s,
        "No figure is a stored answer. If we cannot recompute it from claims and contract rules, it does not belong in a Glass report.",
        title="Reporting promise",
        top=Inches(6.45),
    )

    # 5 Transparency
    s = blank_slide(prs)
    eyebrow(s, "Pillar two")
    title(s, "Transparency — without a single fee model", width=Inches(12))
    lede(
        s,
        "Transparency is the product layer. Pricing is packaging. Sponsors who need PEPM or shared savings still get claim-level derivation.",
        top=Inches(2.15),
    )
    card(
        s,
        Inches(0.7),
        Inches(3.15),
        Inches(3.85),
        Inches(2.7),
        "Always on — Transparency layer",
        [
            "Lesser-of arms on every claim",
            "Source documents with citations",
            "Settlement: invoice vs remit vs rebates",
            "Guarantees from sponsor-visible records",
        ],
    )
    card(
        s,
        Inches(4.75),
        Inches(3.15),
        Inches(3.85),
        Inches(2.7),
        "Your choice — Commercial shell",
        [
            "Full pass-through",
            "Transparent PEPM",
            "Shared savings on verified interventions",
            "Outcomes / SLA dollars at risk",
            "Disclosed hybrid margin if required",
        ],
    )
    card(
        s,
        Inches(8.8),
        Inches(3.15),
        Inches(3.85),
        Inches(2.7),
        "Rule — Disclosure",
        [
            "If Glass earns money, it appears as a named fee, share, credit, or disclosed margin",
            "Never an unexplained plan-vs-pharmacy gap",
        ],
    )
    callout(
        s,
        "“Buy the fee model that fits your RFP. You still get the ledger. You just cannot buy opacity.”",
        title="Line for the room",
        top=Inches(6.2),
    )

    # 6 Service
    s = blank_slide(prs)
    eyebrow(s, "Pillar three")
    title(s, "Excellent service through chat and text", width=Inches(12))
    lede(
        s,
        "Members should not need HR or a 45-minute hold to learn why a claim rejected. Glass answers in chat and SMS from the same rules engine that priced the claim.",
        top=Inches(2.15),
    )
    stat_card(s, Inches(0.7), Inches(3.1), Inches(5.85), Inches(1.15), "Chat", "In-app / web member service agent")
    stat_card(s, Inches(6.75), Inches(3.1), Inches(5.85), Inches(1.15), "Text", "SMS for status, rejects, and next steps")
    table_slide(
        s,
        ["Channel", "Example", "Why it works"],
        [
            ["Chat", "“Is Ozempic covered for me?”", "Reads formulary + member eligibility live"],
            ["Chat", "“What’s the status of my prior auth?”", "Follows into the PA record without a second ask"],
            ["Text", "Reject at the pharmacy → why + what to do", "Pushes the answer when friction happens"],
            ["Text", "PA approved / more info needed / expiring", "Keeps the member out of the call center"],
            ["Either", "Cost share, accumulators, alternatives", "Every number comes from a tool on the book"],
        ],
        top=Inches(4.45),
        col_widths=[Inches(1.6), Inches(5.2), Inches(5.1)],
    )

    # 7 Pricing
    s = blank_slide(prs)
    eyebrow(s, "Commercial menu")
    title(s, "Flexible pricing under the same glass", width=Inches(12))
    lede(
        s,
        "Reporting and service do not change with the fee model. Only how Glass appears on the invoice changes — always disclosed.",
        top=Inches(2.15),
    )
    table_slide(
        s,
        ["Model", "How Glass is paid", "Reporting & service"],
        [
            ["Full pass-through", "Low admin PEPM/PMPM only", "Full ledger + chat/text included"],
            ["Transparent PEPM", "Higher PEPM for ops + service modules", "Same reports; service SLAs explicit"],
            ["Shared savings", "% of verified, reproducible savings", "Savings measured in the same reporting layer"],
            ["Outcomes / SLA at risk", "Base fee + credits if targets miss", "Service and PA turnaround on the scorecard"],
            ["Disclosed hybrid margin", "Named, capped margin on claims", "Margin line visible on claims and invoices"],
        ],
        top=Inches(3.1),
        col_widths=[Inches(3.2), Inches(4.4), Inches(4.3)],
    )
    callout(
        s,
        "No unexplained plan-vs-pharmacy gap. Chat and text answers stay grounded in the book under every commercial model.",
        title="Non-negotiable",
        top=Inches(6.35),
    )

    # 8 Proof
    s = blank_slide(prs)
    eyebrow(s, "Demo path")
    title(s, "Proof points to run live")
    lede(s, "Six stops. End on the channel this sponsor’s members will actually use.", top=Inches(2.05))
    table_slide(
        s,
        ["Stop", "Show", "Pillar"],
        [
            ["1", "Sponsor dashboard → totals match the claim ledger", "Reporting"],
            ["2", "Contract reports + data agent briefing", "Reporting"],
            ["3", "Open a claim: lesser-of arms + source citation", "Transparency"],
            ["4", "Name the fee model they buy — every dollar still labeled", "Transparency"],
            ["5", "Member chat: coverage / PA status with live tools", "Service"],
            ["6", "Text moment: reject or PA update as plain-language SMS", "Service"],
        ],
        top=Inches(2.95),
        col_widths=[Inches(1.2), Inches(8.2), Inches(2.5)],
    )
    card(
        s,
        Inches(0.7),
        Inches(5.85),
        Inches(5.85),
        Inches(1.25),
        "For finance / benefits",
        ["Stay on reports and transparency until they trust the numbers — then show service."],
    )
    card(
        s,
        Inches(6.75),
        Inches(5.85),
        Inches(5.85),
        Inches(1.25),
        "For HR / experience owners",
        ["Lead with chat and text, then show the answer matches the claim."],
    )

    # 9 Ask
    s = blank_slide(prs, dark=True)
    eyebrow(s, "Next step", dark=True)
    title(s, "The ask", dark=True)
    card(
        s,
        Inches(0.7),
        Inches(2.4),
        Inches(5.85),
        Inches(2.6),
        "Near-term engagement",
        [
            "Walk finance through recomputable contract reports",
            "Pick a fee model from the menu — keep full transparency",
            "Run live chat on 5 real member questions from last month",
            "Prototype SMS for rejects + PA status on a pilot population",
        ],
        dark=True,
    )
    card(
        s,
        Inches(6.75),
        Inches(2.4),
        Inches(5.85),
        Inches(2.6),
        "What we need from you",
        [
            "Current reporting gaps finance cannot close",
            "Preferred fee construct for the next RFP",
            "Top member call / HR ticket reasons",
            "Whether SMS is allowed for benefit communications",
        ],
        dark=True,
    )
    close = add_textbox(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.4))
    tf = close.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Reports you can trust. Transparency that survives any fee model. Answers by chat and text."
    set_run(r, 26, False, WHITE)
    r.font.name = "Georgia"

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
